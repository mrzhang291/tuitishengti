#!/usr/bin/env python3
"""Polish a Pandoc-generated math DOCX without changing OMML equations."""

from __future__ import annotations

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

from diagnose_html_math import read_text


IMG_RE = re.compile(r"<img\b[^>]*>", re.IGNORECASE | re.DOTALL)
ATTR_RE = re.compile(r"([\w-]+)=([\"'])(.*?)\2", re.IGNORECASE | re.DOTALL)


def html_image_sizes(html_file: Path) -> list[tuple[int, int]]:
    source = read_text(html_file)
    sizes: list[tuple[int, int]] = []
    for match in IMG_RE.finditer(source):
        attrs = {key.lower(): value for key, _, value in ATTR_RE.findall(match.group(0))}
        width = attrs.get("data-width") or attrs.get("width")
        height = attrs.get("data-height") or attrs.get("height")
        if width and height and width.isdigit() and height.isdigit():
            sizes.append((int(width), int(height)))
    return sizes


def set_east_asian_font(style, font_name: str) -> None:
    style.font.name = font_name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def format_runs(paragraph, font_name: str, size_pt: float, bold: bool | None = None) -> None:
    for run in paragraph.runs:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        run.font.size = Pt(size_pt)
        if bold is not None:
            run.bold = bold


def set_paragraph_spacing(paragraph, before_pt: float = 0, after_pt: float = 2, line_spacing: float = 1.12) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before_pt)
    fmt.space_after = Pt(after_pt)
    fmt.line_spacing = line_spacing


def paragraph_has_picture(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def paragraph_text(paragraph) -> str:
    return re.sub(r"\s+", " ", paragraph.text).strip()


def is_section_heading(text: str) -> bool:
    return bool(re.match(r"^[\u4e00\u4e8c\u4e09\u56db\u4e94\u516d\u4e03\u516b\u4e5d\u5341]+、", text))


def is_question_start(text: str) -> bool:
    return bool(re.match(r"^\d+[．.]", text))


def is_solution_section(text: str) -> bool:
    return is_section_heading(text) and bool(re.search(r"(\u89e3\u7b54\u9898|\u7b80\u7b54\u9898|\u8ba1\u7b97\u9898|\u8bc1\u660e)", text))


def insert_blank_paragraph_after(paragraph, line_height_pt: float):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    new_paragraph.add_run(" ")
    set_paragraph_spacing(new_paragraph, before_pt=0, after_pt=0, line_spacing=1.0)
    new_paragraph.paragraph_format.line_spacing = Pt(line_height_pt)
    return new_paragraph


def add_solution_space(doc: Document, blank_lines: int, line_height_pt: float) -> None:
    if blank_lines <= 0:
        return

    paragraphs = list(doc.paragraphs)
    solution_start = None
    for index, paragraph in enumerate(paragraphs):
        if is_solution_section(paragraph_text(paragraph)):
            solution_start = index
            break
    if solution_start is None:
        return

    question_starts = [
        index
        for index in range(solution_start + 1, len(paragraphs))
        if is_question_start(paragraph_text(paragraphs[index]))
    ]
    if not question_starts:
        return

    insertion_points = []
    for pos, start_index in enumerate(question_starts):
        next_start = question_starts[pos + 1] if pos + 1 < len(question_starts) else len(paragraphs)
        end_index = next_start - 1
        while end_index > start_index and not paragraph_text(paragraphs[end_index]) and not paragraph_has_picture(paragraphs[end_index]):
            end_index -= 1
        insertion_points.append(paragraphs[end_index])

    for paragraph in reversed(insertion_points):
        cursor = paragraph
        for _ in range(blank_lines):
            cursor = insert_blank_paragraph_after(cursor, line_height_pt)


def polish_docx(
    input_docx: Path,
    output_docx: Path,
    html_file: Path | None,
    max_image_width_in: float,
    image_scale: float,
    question_after_pt: float,
    solution_blank_lines: int,
    solution_line_height_pt: float,
) -> None:
    if input_docx.resolve() != output_docx.resolve():
        shutil.copyfile(input_docx, output_docx)

    doc = Document(output_docx)

    for section in doc.sections:
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)

    normal = doc.styles["Normal"]
    set_east_asian_font(normal, "SimSun")
    normal.font.size = Pt(10.5)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        if style_name in doc.styles:
            set_east_asian_font(doc.styles[style_name], "SimSun")

    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph_text(paragraph)
        has_pic = paragraph_has_picture(paragraph)
        set_paragraph_spacing(paragraph)

        if index == 0 and text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_runs(paragraph, "SimSun", 16, True)
            set_paragraph_spacing(paragraph, after_pt=8, line_spacing=1.08)
            continue

        if is_section_heading(text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_runs(paragraph, "SimSun", 11, True)
            set_paragraph_spacing(paragraph, before_pt=6, after_pt=5, line_spacing=1.1)
        elif is_question_start(text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_runs(paragraph, "SimSun", 10.5)
            set_paragraph_spacing(paragraph, before_pt=3, after_pt=question_after_pt, line_spacing=1.14)
        elif has_pic and not text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_spacing(paragraph, before_pt=2, after_pt=2, line_spacing=1.0)
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            format_runs(paragraph, "SimSun", 10.5)

    sizes = html_image_sizes(html_file) if html_file else []
    for index, shape in enumerate(doc.inline_shapes):
        if index < len(sizes):
            width_px, height_px = sizes[index]
            width_in = min((width_px / 96.0) * image_scale, max_image_width_in)
            height_in = width_in * height_px / width_px
        else:
            width_in = min(shape.width / 914400, max_image_width_in)
            height_in = width_in * shape.height / shape.width
        shape.width = Inches(width_in)
        shape.height = Inches(height_in)

    add_solution_space(doc, solution_blank_lines, solution_line_height_pt)
    doc.save(output_docx)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx_file", type=Path)
    parser.add_argument("-o", "--output", type=Path, help="Output DOCX path. Defaults to overwriting input.")
    parser.add_argument("--html", type=Path, help="Original or normalized HTML used to recover image data-width sizes.")
    parser.add_argument("--max-image-width-in", type=float, default=4.2)
    parser.add_argument("--image-scale", type=float, default=1.0, help="Scale factor applied to HTML data-width pixels.")
    parser.add_argument("--question-after-pt", type=float, default=6.0, help="Space after numbered question paragraphs.")
    parser.add_argument("--solution-blank-lines", type=int, default=8, help="Blank writing lines after each solution question.")
    parser.add_argument("--solution-line-height-pt", type=float, default=16.0, help="Height of each inserted solution blank line.")
    args = parser.parse_args()

    output = args.output or args.docx_file
    polish_docx(
        args.docx_file,
        output,
        args.html,
        args.max_image_width_in,
        args.image_scale,
        args.question_after_pt,
        args.solution_blank_lines,
        args.solution_line_height_pt,
    )
    print(f"Polished DOCX written to {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
